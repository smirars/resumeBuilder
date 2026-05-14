import tempfile
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from app.models.resume import Resume


def generate_docx(resume: Resume) -> str:
    doc = Document()
    font = resume.font or "Arial"

    _set_margins(doc)
    _set_default_style(doc, font)

    for block in resume.blocks_data:
        if not block.get("visible", True):
            continue
        bt = block.get("block_type")
        content = block.get("content", {})
        handlers = {
            "contacts": _add_contacts,
            "summary": _add_summary,
            "experience": _add_experience,
            "education": _add_education,
            "skills": _add_skills,
            "languages": _add_languages,
            "projects": _add_projects,
        }
        handler = handlers.get(bt)
        if handler:
            handler(doc, content, font)

    tmp = tempfile.NamedTemporaryFile(suffix=".docx", delete=False)
    tmp.close()
    doc.save(tmp.name)
    return tmp.name


def _set_margins(doc):
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)


def _set_default_style(doc, font):
    style = doc.styles["Normal"]
    style.font.name = font
    style.font.size = Pt(10)


def _heading(doc, text, font):
    p = doc.add_paragraph()
    run = p.add_run(text.upper())
    run.bold = True
    run.font.name = font
    run.font.size = Pt(11)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)


def _para(doc, text, font, bold=False, space_after=4):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.font.name = font
    p.paragraph_format.space_after = Pt(space_after)
    return p


def _add_contacts(doc, content, font):
    p = doc.add_heading(content.get("name", ""), level=1)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in p.runs:
        run.font.name = font

    details = [content.get(f, "") for f in ("email", "phone", "location", "website")]
    details = [d for d in details if d]
    if details:
        p2 = doc.add_paragraph(" | ".join(details))
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.space_after = Pt(6)
        for run in p2.runs:
            run.font.name = font


def _add_summary(doc, content, font):
    _heading(doc, "О себе", font)
    _para(doc, content.get("text", ""), font)


def _add_experience(doc, content, font):
    _heading(doc, "Опыт работы", font)
    for item in content.get("items", []):
        p = doc.add_paragraph()
        r1 = p.add_run(item.get("position", ""))
        r1.bold = True
        r1.font.name = font
        r2 = p.add_run(f"  —  {item.get('company', '')}")
        r2.font.name = font

        end = item.get("end_date") or "н.в."
        _para(doc, f"{item.get('start_date', '')} — {end}", font, space_after=2)

        if item.get("description"):
            _para(doc, item["description"], font)


def _add_education(doc, content, font):
    _heading(doc, "Образование", font)
    for item in content.get("items", []):
        p = doc.add_paragraph()
        r = p.add_run(item.get("institution", ""))
        r.bold = True
        r.font.name = font
        _para(doc, f"{item.get('degree', '')}  —  {item.get('year', '')}", font)


def _add_skills(doc, content, font):
    _heading(doc, "Навыки", font)
    items = content.get("items", [])
    if items:
        _para(doc, ", ".join(items), font)


def _add_languages(doc, content, font):
    _heading(doc, "Языки", font)
    items = content.get("items", [])
    if items:
        _para(doc, ", ".join(items), font)


def _add_projects(doc, content, font):
    _heading(doc, "Проекты", font)
    for item in content.get("items", []):
        p = doc.add_paragraph()
        r = p.add_run(item.get("name", ""))
        r.bold = True
        r.font.name = font
        if item.get("description"):
            _para(doc, item["description"], font)
        if item.get("url"):
            _para(doc, item["url"], font, space_after=6)
